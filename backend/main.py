from fastapi import FastAPI, Depends, HTTPException, Header, File, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from sqlalchemy.orm import Session
from typing import List
from uuid import UUID
from datetime import datetime
from pydantic import BaseModel
import openai
from openai import OpenAI

from database import get_db
from models import User, Idea, Team, ProgressUpdate, Note, VisibilityType, UpdateType, UserRole, SystemSettings, AIInsight
import schemas

app = FastAPI(title="Hackathon Coach Management System API V2")

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=False,
    allow_methods=["*"],
    allow_headers=["*"],
)

@app.on_event("startup")
def on_startup():
    from database import engine
    from sqlalchemy import text
    try:
        with engine.execution_options(isolation_level="AUTOCOMMIT").connect() as conn:
            conn.execute(text("ALTER TYPE user_role ADD VALUE IF NOT EXISTS 'DEDICATED_COACH'"))
            conn.execute(text("ALTER TYPE user_role ADD VALUE IF NOT EXISTS 'PARTICIPANT'"))
    except Exception as e:
        print(f"Startup DB patch error (ignored): {e}")

class LoginRequest(BaseModel):
    username: str
    password: str

class StatusUpdateRequest(BaseModel):
    is_active: bool

# --- Authentication ---
def get_current_user(authorization: str = Header(None), db: Session = Depends(get_db)):
    if not authorization:
        raise HTTPException(status_code=401, detail="Missing Authorization Header")
    user_id = authorization.replace("Bearer ", "").strip()
    try:
        user = db.query(User).filter(User.id == UUID(user_id)).first()
        if not user:
            raise HTTPException(status_code=404, detail="User not found")
        if not user.is_active:
            raise HTTPException(status_code=403, detail="Account pending approval or disabled")
        return user
    except ValueError:
        raise HTTPException(status_code=400, detail="Invalid token format")

def get_superadmin(current_user: User = Depends(get_current_user)):
    if current_user.system_role != UserRole.ADMIN:
        raise HTTPException(status_code=403, detail="Superadmin access required")
    return current_user

# --- Auth & Users ---
@app.post("/api/v1/auth/register", response_model=schemas.UserResponse)
def register(user: schemas.UserCreate, db: Session = Depends(get_db)):
    existing = db.query(User).filter(User.username == user.username).first()
    if existing:
        raise HTTPException(status_code=400, detail="Username already exists")
    
    db_user = User(
        username=user.username,
        hashed_password=user.password + "_hashed",
        system_role=user.system_role,
        first_name=user.first_name,
        last_name=user.last_name,
        is_active=False,
        requires_password_change=True
    )
    db.add(db_user)
    db.commit()
    db.refresh(db_user)
    return db_user

@app.post("/api/v1/auth/login")
def login(req: LoginRequest, db: Session = Depends(get_db)):
    user = db.query(User).filter(User.username == req.username).first()
    if not user or user.hashed_password != req.password + "_hashed":
        raise HTTPException(status_code=401, detail="Invalid credentials")
    if not user.is_active:
        raise HTTPException(status_code=403, detail="Account pending approval or disabled")
    return {
        "token": str(user.id),
        "requires_password_change": user.requires_password_change,
        "role": user.system_role
    }

@app.post("/api/v1/auth/change-password")
def change_password(req: schemas.PasswordChangeRequest, current_user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    if current_user.hashed_password != req.old_password + "_hashed":
        raise HTTPException(status_code=400, detail="Incorrect old password")
    current_user.hashed_password = req.new_password + "_hashed"
    current_user.requires_password_change = False
    db.commit()
    return {"message": "Password changed successfully"}

# --- Superadmin Endpoints ---
@app.get("/api/v1/admin/users/pending", response_model=List[schemas.UserResponse])
def get_pending_users(admin: User = Depends(get_superadmin), db: Session = Depends(get_db)):
    return db.query(User).filter(User.is_active == False).all()

@app.put("/api/v1/admin/users/{user_id}/status")
def update_user_status(user_id: UUID, req: StatusUpdateRequest, admin: User = Depends(get_superadmin), db: Session = Depends(get_db)):
    user = db.query(User).filter(User.id == user_id).first()
    if not user:
        raise HTTPException(status_code=404, detail="User not found")
    user.is_active = req.is_active
    db.commit()
    return {"message": "User status updated", "is_active": user.is_active}

@app.post("/api/v1/admin/settings")
def update_settings(req: schemas.SystemSettingsUpdate, admin: User = Depends(get_superadmin), db: Session = Depends(get_db)):
    settings = db.query(SystemSettings).first()
    if not settings:
        settings = SystemSettings()
        db.add(settings)
    if req.openai_api_key is not None:
        settings.openai_api_key = req.openai_api_key
        settings.openai_api_key_added_by = admin.id
    if req.reporting_interval_hours is not None:
        settings.reporting_interval_hours = req.reporting_interval_hours
    db.commit()
    return {"message": "Settings updated"}

@app.get("/api/v1/admin/settings", response_model=schemas.SystemSettingsResponse)
def get_settings(admin: User = Depends(get_superadmin), db: Session = Depends(get_db)):
    settings = db.query(SystemSettings).first()
    has_api_key = bool(settings and settings.openai_api_key)
    interval = settings.reporting_interval_hours if settings else 24
    return schemas.SystemSettingsResponse(has_api_key=has_api_key, reporting_interval_hours=interval)

@app.get("/api/v1/users", response_model=List[schemas.UserResponse])
def get_users(db: Session = Depends(get_db)):
    return db.query(User).all()

@app.get("/api/v1/users/me/alerts")
def get_user_alerts(current_user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    if current_user.system_role == UserRole.ADMIN:
        return []
        
    settings = db.query(SystemSettings).first()
    interval_hours = settings.reporting_interval_hours if settings else 24
    from datetime import timedelta
    interval_td = timedelta(hours=interval_hours)
    
    alerts = []
    now = datetime.utcnow()
    
    if current_user.system_role == UserRole.DEDICATED_COACH:
        if current_user.team_id:
            team = db.query(Team).filter(Team.id == current_user.team_id).first()
            last_note = db.query(Note).filter(
                Note.author_id == current_user.id,
                Note.team_id == current_user.team_id
            ).order_by(Note.created_at.desc()).first()
            
            if not last_note or (now - last_note.created_at) > interval_td:
                alerts.append(f"Prompt: You haven't reported on your team ({team.name}) in over {interval_hours} hours. Please submit a new note!")
                
    elif current_user.system_role == UserRole.COACH:
        teams = db.query(Team).all()
        for t in teams:
            last_note = db.query(Note).filter(
                Note.author_id == current_user.id,
                Note.team_id == t.id
            ).order_by(Note.created_at.desc()).first()
            
            if not last_note or (now - last_note.created_at) > interval_td:
                alerts.append(f"Prompt: You are due for a report on {t.name} (over {interval_hours} hours since your last note).")
                
    return alerts

# --- Ideas & Teams ---
@app.post("/api/v1/ideas", response_model=schemas.IdeaResponse)
def create_idea(idea: schemas.IdeaCreate, current_user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    db_idea = Idea(**idea.model_dump(), created_by=current_user.id)
    db.add(db_idea)
    db.commit()
    db.refresh(db_idea)
    return db_idea

@app.get("/api/v1/ideas", response_model=List[schemas.IdeaResponse])
def get_ideas(db: Session = Depends(get_db)):
    return db.query(Idea).all()

@app.post("/api/v1/teams", response_model=schemas.TeamResponse)
def create_team(team_req: schemas.TeamCreate, current_user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    if current_user.system_role != UserRole.ADMIN:
        raise HTTPException(status_code=403, detail="Only Admins can create teams")
        
    db_idea = Idea(**team_req.idea.model_dump(), created_by=current_user.id)
    db.add(db_idea)
    db.flush() # get idea id
    
    db_team = Team(name=team_req.name, coach_id=team_req.coach_id, idea_id=db_idea.id)
    db.add(db_team)
    db.flush() # get team id
    
    # Assign the team_id to the dedicated coach so they have access
    coach_user = db.query(User).filter(User.id == team_req.coach_id).first()
    if coach_user:
        coach_user.team_id = db_team.id
        
    db.commit()
    db.refresh(db_team)
    return db_team

@app.post("/api/v1/teams/{team_id}/presentation")
def upload_presentation(team_id: UUID, file: UploadFile = File(...), current_user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    if current_user.system_role != UserRole.ADMIN:
        raise HTTPException(status_code=403, detail="Only Admins can upload presentations")
        
    team = db.query(Team).filter(Team.id == team_id).first()
    if not team or not team.idea_id:
        raise HTTPException(status_code=404, detail="Team or Idea not found")
        
    idea = db.query(Idea).filter(Idea.id == team.idea_id).first()
    
    import io
    from pptx import Presentation
    
    try:
        content = file.file.read()
        prs = Presentation(io.BytesIO(content))
        text_runs = []
        for slide in prs.slides:
            for shape in slide.shapes:
                if hasattr(shape, "text"):
                    text_runs.append(shape.text)
        
        extracted_text = "\n".join(text_runs)
        idea.presentation_text = extracted_text
        db.commit()
        
        return {"message": "Presentation uploaded and parsed successfully", "extracted_length": len(extracted_text)}
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Failed to parse PPTX file: {str(e)}")

@app.get("/api/v1/teams", response_model=List[schemas.TeamResponse])
def get_teams(current_user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    if current_user.system_role == UserRole.DEDICATED_COACH:
        if current_user.team_id:
            return db.query(Team).filter(Team.id == current_user.team_id).all()
        return []
    return db.query(Team).all()

@app.post("/api/v1/teams/{team_id}/participants", response_model=schemas.UserResponse)
def add_participant(team_id: UUID, participant: schemas.ParticipantCreate, current_user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    # Check permissions (must be Admin, Global Coach, or the Dedicated Coach for this team)
    if current_user.system_role not in [UserRole.ADMIN, UserRole.COACH]:
        if current_user.system_role != UserRole.DEDICATED_COACH or current_user.team_id != team_id:
            raise HTTPException(status_code=403, detail="Not authorized to add members to this team")
    
    # Check if username exists
    if db.query(User).filter(User.username == participant.username).first():
        raise HTTPException(status_code=400, detail="Username already exists")

    # Create the user for the participant
    db_user = User(
        username=participant.username,
        hashed_password=participant.password + "_hashed",
        system_role=UserRole.PARTICIPANT,
        is_active=True,
        requires_password_change=True,
        first_name=participant.first_name,
        last_name=participant.last_name,
        team_id=team_id,
        team_role=participant.team_role
    )
    db.add(db_user)
    db.commit()
    db.refresh(db_user)
    return db_user

# --- Progress Updates & Notes ---
@app.post("/api/v1/progress_updates", response_model=schemas.ProgressUpdateResponse)
def create_progress_update(update: schemas.ProgressUpdateCreate, db: Session = Depends(get_db)):
    db_update = ProgressUpdate(**update.model_dump())
    db.add(db_update)
    db.commit()
    db.refresh(db_update)
    return db_update

@app.post("/api/v1/notes", response_model=schemas.NoteResponse)
def create_note(note: schemas.NoteCreate, current_user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    if not note.team_id and not note.progress_update_id and not note.role_tag:
        raise HTTPException(status_code=400, detail="Note must be linked to a team, role, or progress update")
    
    db_note = Note(**note.model_dump(), author_id=current_user.id)
    db.add(db_note)
    db.commit()
    db.refresh(db_note)
    return db_note

# --- AI Insights ---
@app.get("/api/v1/teams/{team_id}/insights", response_model=schemas.AIInsightResponse)
def get_team_insights(team_id: UUID, current_user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    if current_user.system_role not in [UserRole.ADMIN, UserRole.COACH, UserRole.DEDICATED_COACH]:
        raise HTTPException(status_code=403, detail="Not authorized to generate insights")
        
    settings = db.query(SystemSettings).first()
    if not settings or not settings.openai_api_key:
        raise HTTPException(status_code=400, detail="OpenAI API key not configured by Superadmin")

    notes = db.query(Note).filter(Note.team_id == team_id).all()
    
    latest_insight = db.query(AIInsight).filter(AIInsight.team_id == team_id).order_by(AIInsight.created_at.desc()).first()
    latest_note = db.query(Note).filter(Note.team_id == team_id).order_by(Note.created_at.desc()).first()
    
    # Cache hit logic: Only generate new insight if there is a note newer than the latest insight
    if latest_insight:
        if not latest_note or latest_insight.created_at >= latest_note.created_at:
            return latest_insight

    try:
        team = db.query(Team).filter(Team.id == team_id).first()
        idea = db.query(Idea).filter(Idea.id == team.idea_id).first() if team else None
        
        client = OpenAI(api_key=settings.openai_api_key)
        prompt = f"Team: {team.name if team else 'Unknown'}\n"
        if idea:
            prompt += f"Project Idea: {idea.title}\n"
            prompt += f"Problem Statement: {idea.problem_statement}\n"
            prompt += f"Proposed Solution: {idea.proposed_solution}\n"
            if idea.presentation_text:
                prompt += f"Presentation Text Context: {idea.presentation_text}\n"
        
        prompt += f"\nReview the following notes for this team and provide actionable insight on how the coach should guide this team/individual:\n"
        for n in notes:
            prompt += f"- Note ({n.visibility}): {n.content}\n"
            
        response = client.chat.completions.create(
            model="gpt-3.5-turbo",
            messages=[
                {"role": "system", "content": "You are an expert hackathon coach advisor. Keep insights brief, actionable, and encouraging."},
                {"role": "user", "content": prompt}
            ]
        )
        insight_content = response.choices[0].message.content
        new_insight = AIInsight(team_id=team_id, content=insight_content)
        db.add(new_insight)
        db.commit()
        db.refresh(new_insight)
        return new_insight
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"OpenAI Error: {str(e)}")

# --- Coach Snapshot ---
@app.get("/api/v1/teams/{team_id}/snapshot", response_model=schemas.TeamSnapshotResponse)
def get_team_snapshot(team_id: UUID, current_user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    if current_user.system_role == UserRole.DEDICATED_COACH and current_user.team_id != team_id:
        raise HTTPException(status_code=403, detail="Not authorized to view this team")
        
    team = db.query(Team).filter(Team.id == team_id).first()
    if not team:
        raise HTTPException(status_code=404, detail="Team not found")
        
    members = db.query(User).filter(User.team_id == team_id).all()
    idea = db.query(Idea).filter(Idea.id == team.idea_id).first() if team.idea_id else None
    
    updates = db.query(ProgressUpdate).filter(ProgressUpdate.team_id == team_id).all()
    team_milestones = [u for u in updates if u.type == UpdateType.TEAM_MILESTONE]
    individual_updates = [u for u in updates if u.type == UpdateType.INDIVIDUAL_TASK]
    
    all_notes = db.query(Note).filter(
        (Note.team_id == team_id) | 
        (Note.progress_update_id.in_([u.id for u in updates] if updates else []))
    ).all()
    
    visible_notes = []
    for note in all_notes:
        if note.visibility == VisibilityType.PUBLIC:
            if current_user.system_role in [UserRole.ADMIN, UserRole.COACH]:
                visible_notes.append(note)
            elif current_user.system_role == UserRole.DEDICATED_COACH and current_user.team_id == team_id:
                visible_notes.append(note)
        elif note.author_id == current_user.id:
            visible_notes.append(note)

    snapshot_members = [
        schemas.SnapshotTeamMember(id=m.id, name=f"{m.first_name} {m.last_name}", team_role=m.team_role)
        for m in members
    ]

    snapshot_team = schemas.SnapshotTeamInfo(
        id=team.id,
        name=team.name,
        coach_id=team.coach_id,
        idea=idea,
        members=snapshot_members
    )
    
    # Fetch all insights to show history
    all_insights = db.query(AIInsight).filter(AIInsight.team_id == team_id).order_by(AIInsight.created_at.desc()).all()
    
    return schemas.TeamSnapshotResponse(
        team=snapshot_team,
        team_milestones=team_milestones,
        individual_updates=individual_updates,
        notes=visible_notes,
        ai_insights=all_insights
    )

@app.get("/health")
def health_check():
    return {"status": "healthy"}
from fastapi import FastAPI, Depends, HTTPException
from sqlalchemy.orm import Session
from uuid import UUID
from database import get_db
from models import User, Team, Idea, Note, ProgressUpdate, AIInsight, UserRole

def add_delete_endpoints(app: FastAPI):
    def get_superadmin_local(current_user: User):
        if current_user.system_role != UserRole.ADMIN:
            raise HTTPException(status_code=403, detail="Superadmin access required")
        return current_user

    from main import get_current_user

    @app.delete("/api/v1/teams/{team_id}")
    def delete_team(team_id: UUID, current_user: User = Depends(get_current_user), db: Session = Depends(get_db)):
        if current_user.system_role != UserRole.ADMIN:
            raise HTTPException(status_code=403, detail="Superadmin access required")
            
        team = db.query(Team).filter(Team.id == team_id).first()
        if not team:
            raise HTTPException(status_code=404, detail="Team not found")
            
        # Delete related child records explicitly to avoid FK errors without CASCADE set in DB
        db.query(Note).filter(Note.team_id == team_id).delete(synchronize_session=False)
        db.query(ProgressUpdate).filter(ProgressUpdate.team_id == team_id).delete(synchronize_session=False)
        db.query(AIInsight).filter(AIInsight.team_id == team_id).delete(synchronize_session=False)
        
        # Clear coach team_ids
        coaches = db.query(User).filter(User.team_id == team_id).all()
        for c in coaches:
            c.team_id = None
            
        # Delete idea
        if team.idea_id:
            db.query(Idea).filter(Idea.id == team.idea_id).delete(synchronize_session=False)
            
        db.delete(team)
        db.commit()
        return {"message": "Team deleted successfully"}

    @app.delete("/api/v1/users/{user_id}")
    def delete_user(user_id: UUID, current_user: User = Depends(get_current_user), db: Session = Depends(get_db)):
        if current_user.system_role != UserRole.ADMIN:
            raise HTTPException(status_code=403, detail="Superadmin access required")
            
        user = db.query(User).filter(User.id == user_id).first()
        if not user:
            raise HTTPException(status_code=404, detail="User not found")
        if user.system_role == UserRole.ADMIN:
            raise HTTPException(status_code=400, detail="Cannot delete superadmin")
            
        # Delete notes authored by them
        db.query(Note).filter(Note.author_id == user_id).delete(synchronize_session=False)
        
        # Delete user
        db.delete(user)
        db.commit()
        return {"message": "User deleted successfully"}
add_delete_endpoints(app)
